"""Build public/downloads/candron-rfq-bid-scorecard.xlsx.

Usage (from the candron-next root):  python3 scripts/resources/build_rfq_scorecard.py

Four sheets: How to use · RFQ (template to send to bidders) · Bid Comparison (weighted
scorecard with live formulas) · Clarifications (question log). Criteria and weights live in
CRITERIA below and must match lib/resources/rfq-scorecard.ts, which the web page reads.
"""
from openpyxl import Workbook
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

OUT = "public/downloads/candron-rfq-bid-scorecard.xlsx"

NAVY, BLUE, PALE, LINE, GREY, INPUT = "05091F", "0047FF", "EEF2FA", "D5DEEE", "4F5D7A", "FFF8E1"
F = "Arial"
thin = Side(style="thin", color=LINE)
BOX = Border(left=thin, right=thin, top=thin, bottom=thin)
WRAP = Alignment(wrap_text=True, vertical="top")
CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)

# (criterion, weight %, how it is scored, auto formula kind or None for manual 1-5)
CRITERIA = [
    ("Price", 35, "Lowest total price scores 5; others pro rata", "low_price"),
    ("Lead time", 25, "Shortest lead time scores 5; others pro rata", "low_lead"),
    ("Technical compliance", 15, "1–5: meets the specification with no technical deviations", None),
    ("Exceptions & commercial terms", 10, "1–5: fewer exceptions and clearer terms score higher", None),
    ("Warranty & support", 5, "Longest warranty scores 5; others pro rata", "high_warranty"),
    ("Testing & documentation", 5, "1–5: witnessed FAT, test reports, drawings and manuals included", None),
    ("Supplier track record", 5, "1–5: references, past performance, local service", None),
]
BIDDERS = 5
EXAMPLE = [  # name, price, lead, drawings, warranty, freight, fat, exceptions, origin, validity, manual scores
    ("Bidder A (example)", 412000, 38, 3, 24, "Y", "Y", 2, "Canada", 30, {"Technical compliance": 5, "Exceptions & commercial terms": 4, "Testing & documentation": 5, "Supplier track record": 4}),
    ("Bidder B (example)", 389000, 56, 5, 18, "N", "Y", 5, "Imported", 45, {"Technical compliance": 4, "Exceptions & commercial terms": 3, "Testing & documentation": 4, "Supplier track record": 4}),
    ("Bidder C (example)", 445000, 30, 2, 24, "Y", "Y", 1, "Canada", 30, {"Technical compliance": 5, "Exceptions & commercial terms": 5, "Testing & documentation": 5, "Supplier track record": 3}),
]


def title(ws, text, sub):
    ws["A1"] = text
    ws["A1"].font = Font(name=F, size=18, bold=True, color=NAVY)
    ws["A2"] = sub
    ws["A2"].font = Font(name=F, size=10, color=GREY)
    ws.row_dimensions[1].height = 28


def band(ws, row, text, cols):
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=cols)
    c = ws.cell(row=row, column=1, value=text)
    c.font = Font(name=F, size=11, bold=True, color="FFFFFF")
    c.fill = PatternFill("solid", fgColor=NAVY)
    c.alignment = Alignment(vertical="center", indent=1)
    ws.row_dimensions[row].height = 22


def head(c, text):
    c.value = text
    c.font = Font(name=F, size=9, bold=True, color=NAVY)
    c.fill = PatternFill("solid", fgColor=PALE)
    c.alignment = CENTER
    c.border = BOX


def cell(ws, row, col, value=None, bold=False, fill=None, fmt=None, align=WRAP, color=NAVY):
    c = ws.cell(row=row, column=col, value=value)
    c.font = Font(name=F, size=10, bold=bold, color=color)
    c.border = BOX
    c.alignment = align
    if fill:
        c.fill = PatternFill("solid", fgColor=fill)
    if fmt:
        c.number_format = fmt
    return c


def footer(ws, row, cols):
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=cols)
    c = ws.cell(row=row, column=1, value="Free resource from CANDRON Energy Inc. · providers of Canadian-made switchgear, switchboards and transformers · candron.com · +1 (647) 616-2595")
    c.font = Font(name=F, size=8, italic=True, color=GREY)


def setup_print(ws, landscape=True):
    ws.page_setup.orientation = "landscape" if landscape else "portrait"
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.print_options.horizontalCentered = True
    ws.page_margins.left = ws.page_margins.right = 0.4


wb = Workbook()

# ── How to use ───────────────────────────────────────────────────────────────
ws = wb.active
ws.title = "How to use"
ws.sheet_view.showGridLines = False
ws.column_dimensions["A"].width = 4
ws.column_dimensions["B"].width = 96
title(ws, "Electrical Equipment RFQ & Bid Scorecard", "Get quotes you can compare, then score them the same way every time.")
steps = [
    ("1. Fill in the RFQ sheet", "Project details, the equipment list and your dates. Send it to every bidder with your specification and drawings."),
    ("2. Ask every bidder for the same things", "The RFQ lists what each quote must state: price basis, lead time from approved drawings, tests, freight, warranty and exceptions. That is what makes quotes comparable."),
    ("3. Log questions on the Clarifications sheet", "Send the same answer to every bidder so nobody gains an advantage, and note any impact on price or lead time."),
    ("4. Enter the quotes on the Bid Comparison sheet", "Yellow cells are inputs. Price, lead time and warranty score automatically; give the other criteria a 1–5 score."),
    ("5. Check the weights", "Weights must add up to 100%. Change them to suit the project: a schedule-critical job may weight lead time above price."),
    ("6. Read the result", "Each bidder gets a weighted score out of 100 and a rank. Use it to support the decision, not replace judgement on technical risk."),
]
r = 4
for h, t in steps:
    ws.cell(row=r, column=2, value=h).font = Font(name=F, size=11, bold=True, color=NAVY)
    c = ws.cell(row=r + 1, column=2, value=t)
    c.font = Font(name=F, size=10, color=GREY)
    c.alignment = WRAP
    ws.row_dimensions[r + 1].height = 30
    r += 3
ws.cell(row=r, column=2, value="Example data: three example bidders are filled in on the Bid Comparison sheet so you can see how scoring works. Overwrite or clear them before use.").font = Font(name=F, size=10, bold=True, color=BLUE)
r += 2
ws.cell(row=r, column=2, value="This workbook is a decision aid. Check every quote against your specification, and follow your organization’s procurement policy.").font = Font(name=F, size=9, italic=True, color=GREY)
footer(ws, r + 2, 2)

# ── RFQ ──────────────────────────────────────────────────────────────────────
ws = wb.create_sheet("RFQ")
ws.sheet_view.showGridLines = False
widths = [6, 22, 26, 10, 30, 22, 20, 28]
for i, w in enumerate(widths, 1):
    ws.column_dimensions[get_column_letter(i)].width = w
title(ws, "Request for Quotation: Electrical Equipment", "Send this sheet with your specification and drawings. Ask each bidder to answer every line in section C.")
band(ws, 4, "A · Project & contacts", 8)
fields = [("Project name", "Location (city, province)"), ("Owner", "Engineer of record"), ("Your company", "Contact name"),
          ("Email", "Phone"), ("RFQ issue date", "Quotes due"), ("Required on-site date", "Target energization date")]
r = 5
for a, b in fields:
    cell(ws, r, 1, None); ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=2)
    ws.cell(row=r, column=1, value=a).font = Font(name=F, size=9, bold=True, color=GREY)
    ws.merge_cells(start_row=r, start_column=3, end_row=r, end_column=4)
    cell(ws, r, 3, fill=INPUT)
    ws.cell(row=r, column=5, value=b).font = Font(name=F, size=9, bold=True, color=GREY)
    ws.cell(row=r, column=5).border = BOX
    ws.merge_cells(start_row=r, start_column=6, end_row=r, end_column=8)
    cell(ws, r, 6, fill=INPUT)
    ws.row_dimensions[r].height = 20
    r += 1
r += 1
band(ws, r, "B · Equipment scope", 8)
r += 1
for i, h in enumerate(["Item", "Tag", "Equipment type", "Qty", "Key ratings (V, A or kVA, kA, BIL)", "Enclosure / location", "Standards", "Notes"], 1):
    head(ws.cell(row=r, column=i), h)
ws.row_dimensions[r].height = 30
for k in range(8):
    r += 1
    cell(ws, r, 1, k + 1, align=CENTER)
    for col in range(2, 9):
        cell(ws, r, col, fill=INPUT)
    ws.row_dimensions[r].height = 20
r += 2
band(ws, r, "C · Every quote must state", 8)
r += 1
for i, h in enumerate(["#", "Item", "", "", "Bidder’s answer", "", "", ""], 1):
    head(ws.cell(row=r, column=i), h)
ws.merge_cells(start_row=r, start_column=2, end_row=r, end_column=4)
ws.merge_cells(start_row=r, start_column=5, end_row=r, end_column=8)
musts = [
    "Unit and total price, in CAD, taxes excluded",
    "Price validity (days) and any escalation clause (e.g. metals)",
    "Lead time in weeks, counted from approved drawings",
    "Submittal drawing turnaround after purchase order (weeks)",
    "Tests included: routine tests, factory acceptance test, witness option",
    "Freight and delivery terms, including offloading and site access",
    "Warranty period and when it starts (shipment or energization)",
    "Certification: CSA, cUL or field evaluation; utility approvals if needed",
    "Recommended spare parts and their prices",
    "Start-up, commissioning or field service support offered",
    "Country of manufacture and final assembly",
    "Payment terms",
    "Exceptions, deviations and clarifications to the specification",
]
for k, m in enumerate(musts, 1):
    r += 1
    cell(ws, r, 1, k, align=CENTER)
    ws.merge_cells(start_row=r, start_column=2, end_row=r, end_column=4)
    cell(ws, r, 2, m)
    for col in (3, 4):
        ws.cell(row=r, column=col).border = BOX
    ws.merge_cells(start_row=r, start_column=5, end_row=r, end_column=8)
    cell(ws, r, 5, fill=INPUT)
    for col in range(6, 9):
        ws.cell(row=r, column=col).border = BOX
    ws.row_dimensions[r].height = 28
r += 2
band(ws, r, "D · Attachments sent with this RFQ", 8)
for a in ["Single-line diagram", "Specification sections", "Layout / site drawings", "Utility or authority requirements"]:
    r += 1
    ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=8)
    cell(ws, r, 1, f"☐  {a}")
footer(ws, r + 2, 8)
setup_print(ws)

# ── Bid Comparison ───────────────────────────────────────────────────────────
ws = wb.create_sheet("Bid Comparison")
ws.sheet_view.showGridLines = False
ws.column_dimensions["A"].width = 34
ws.column_dimensions["B"].width = 10
ws.column_dimensions["C"].width = 40
for k in range(BIDDERS):
    ws.column_dimensions[get_column_letter(4 + k)].width = 17
title(ws, "Bid Comparison", "Yellow cells are inputs. Price, lead time and warranty score automatically; give the rest a score from 1 to 5.")
first, last = get_column_letter(4), get_column_letter(3 + BIDDERS)

band(ws, 4, "1 · Quote details", 3 + BIDDERS)
rows_in = [("Bidder", "text"), ("Total price (CAD, before tax)", "money"), ("Lead time (weeks from approved drawings)", "num"),
           ("Drawing turnaround (weeks)", "num"), ("Warranty (months)", "num"), ("Freight & offload included? (Y/N)", "yn"),
           ("Witnessed FAT included? (Y/N)", "yn"), ("Number of exceptions to the spec", "num"), ("Country of manufacture", "text"),
           ("Price validity (days)", "num")]
r = 5
head(ws.cell(row=r, column=1), "Quote detail")
ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=3)
for k in range(BIDDERS):
    head(ws.cell(row=r, column=4 + k), f"Bidder {k + 1}")
ROW = {}
yn = DataValidation(type="list", formula1='"Y,N"', allow_blank=True)
ws.add_data_validation(yn)
for label, kind in rows_in:
    r += 1
    ROW[label] = r
    ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=3)
    cell(ws, r, 1, label, bold=True)
    for col in (2, 3):
        ws.cell(row=r, column=col).border = BOX
    for k in range(BIDDERS):
        ex = EXAMPLE[k] if k < len(EXAMPLE) else None
        val = None
        if ex:
            val = {"Bidder": ex[0], "Total price (CAD, before tax)": ex[1], "Lead time (weeks from approved drawings)": ex[2],
                   "Drawing turnaround (weeks)": ex[3], "Warranty (months)": ex[4], "Freight & offload included? (Y/N)": ex[5],
                   "Witnessed FAT included? (Y/N)": ex[6], "Number of exceptions to the spec": ex[7], "Country of manufacture": ex[8],
                   "Price validity (days)": ex[9]}[label]
        fmt = '"$"#,##0' if kind == "money" else ("0" if kind == "num" else None)
        c = cell(ws, r, 4 + k, val, fill=INPUT, fmt=fmt, align=CENTER, bold=(label == "Bidder"))
        if kind == "yn":
            yn.add(c)

r += 2
band(ws, r, "2 · Weighted scoring", 3 + BIDDERS)
r += 1
head(ws.cell(row=r, column=1), "Criterion")
head(ws.cell(row=r, column=2), "Weight")
head(ws.cell(row=r, column=3), "How it is scored")
for k in range(BIDDERS):
    head(ws.cell(row=r, column=4 + k), f"Score 1–5 · Bidder {k + 1}")
ws.row_dimensions[r].height = 30
score_first = r + 1
manual = DataValidation(type="whole", operator="between", formula1="1", formula2="5", allow_blank=True,
                        showErrorMessage=True, errorTitle="Score 1 to 5", error="Enter a whole number from 1 to 5.")
ws.add_data_validation(manual)
price_r, lead_r, warr_r = ROW["Total price (CAD, before tax)"], ROW["Lead time (weeks from approved drawings)"], ROW["Warranty (months)"]
for name, weight, how, kind in CRITERIA:
    r += 1
    cell(ws, r, 1, name, bold=True)
    cell(ws, r, 2, weight / 100, fill=INPUT, fmt="0%", align=CENTER)
    cell(ws, r, 3, how, color=GREY)
    ws.row_dimensions[r].height = 30
    for k in range(BIDDERS):
        col = get_column_letter(4 + k)
        if kind == "low_price":
            f = f'=IF({col}{price_r}>0,ROUND(MIN(${first}${price_r}:${last}${price_r})/{col}{price_r}*5,2),"")'
        elif kind == "low_lead":
            f = f'=IF({col}{lead_r}>0,ROUND(MIN(${first}${lead_r}:${last}${lead_r})/{col}{lead_r}*5,2),"")'
        elif kind == "high_warranty":
            f = f'=IF({col}{warr_r}>0,ROUND({col}{warr_r}/MAX(${first}${warr_r}:${last}${warr_r})*5,2),"")'
        else:
            ex = EXAMPLE[k] if k < len(EXAMPLE) else None
            f = ex[10].get(name) if ex else None
        c = cell(ws, r, 4 + k, f, fill=None if kind else INPUT, fmt="0.00" if kind else "0", align=CENTER)
        if not kind:
            manual.add(c)
score_last = r
r += 1
cell(ws, r, 1, "Weights total (must be 100%)", bold=True)
cell(ws, r, 2, f"=SUM(B{score_first}:B{score_last})", fmt="0%", align=CENTER, bold=True)
cell(ws, r, 3, f'=IF(ROUND(B{r},4)=1,"OK","Adjust weights to total 100%")', color=GREY)
weights_r = r
ws.conditional_formatting.add(f"B{r}", FormulaRule(formula=[f"ROUND(B{r},4)<>1"], fill=PatternFill("solid", fgColor="FDE8E8"), font=Font(color="B42318", bold=True)))
r += 2
band(ws, r, "3 · Result", 3 + BIDDERS)
r += 1
total_r = r
cell(ws, r, 1, "Weighted score (out of 100)", bold=True)
ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=3)
for k in range(BIDDERS):
    col = get_column_letter(4 + k)
    f = (f'=IF(COUNT({col}{score_first}:{col}{score_last})=0,"",ROUND(SUMPRODUCT($B${score_first}:$B${score_last},'
         f'{col}{score_first}:{col}{score_last})/5*100,1))')
    cell(ws, r, 4 + k, f, fmt="0.0", align=CENTER, bold=True)
ws.conditional_formatting.add(f"{first}{r}:{last}{r}", FormulaRule(formula=[f'AND({first}{r}<>"",{first}{r}=MAX(${first}${r}:${last}${r}))'],
                              fill=PatternFill("solid", fgColor="E3F5EC"), font=Font(color="11804B", bold=True)))
r += 1
cell(ws, r, 1, "Rank", bold=True)
ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=3)
for k in range(BIDDERS):
    col = get_column_letter(4 + k)
    cell(ws, r, 4 + k, f'=IF({col}{total_r}="","",RANK({col}{total_r},${first}${total_r}:${last}${total_r}))', fmt="0", align=CENTER, bold=True)
r += 1
cell(ws, r, 1, "Notes on the decision", bold=True)
ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=3)
ws.merge_cells(start_row=r, start_column=4, end_row=r, end_column=3 + BIDDERS)
cell(ws, r, 4, fill=INPUT)
ws.row_dimensions[r].height = 46
ws.freeze_panes = "D6"
footer(ws, r + 2, 3 + BIDDERS)
setup_print(ws)

# ── Clarifications ───────────────────────────────────────────────────────────
ws = wb.create_sheet("Clarifications")
ws.sheet_view.showGridLines = False
for i, w in enumerate([6, 18, 44, 13, 44, 13, 30], 1):
    ws.column_dimensions[get_column_letter(i)].width = w
title(ws, "Clarifications Log", "Send every answer to every bidder. Note any change to price or lead time.")
r = 4
for i, h in enumerate(["#", "Bidder", "Question", "Date asked", "Answer (sent to all bidders)", "Date answered", "Impact on price / lead time"], 1):
    head(ws.cell(row=r, column=i), h)
ws.row_dimensions[r].height = 30
for k in range(15):
    r += 1
    cell(ws, r, 1, k + 1, align=CENTER)
    for col in range(2, 8):
        cell(ws, r, col, fill=INPUT, fmt="yyyy-mm-dd" if col in (4, 6) else None)
    ws.row_dimensions[r].height = 22
ws.freeze_panes = "A5"
footer(ws, r + 2, 7)
setup_print(ws)

wb.active = 0
wb.save(OUT)
print("wrote", OUT)
