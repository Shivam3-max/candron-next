"""Build the Division 26 guide specification sections as editable Word files.

Usage:  python3 scripts/resources/build_div26_specs.py
Writes downloads/specs/*.docx (served by app/downloads/[...path]/route.ts). Section list and summaries for the web page live in
lib/resources/div26.ts; keep the file names in step.
"""
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor, Cm

OUT = Path(__file__).resolve().parents[2] / "downloads" / "specs"
NAVY = RGBColor(0x05, 0x09, 0x1F)
BLUE = RGBColor(0x00, 0x47, 0xFF)
GREY = RGBColor(0x5B, 0x66, 0x80)

COMMON_INTRO = (
    "Guide specification for editing by the engineer of record. Text in [brackets] is an option or a "
    "value to fill in; delete what does not apply. Shaded notes are for the specifier and should be "
    "deleted before issue. Confirm the current edition of every referenced standard. Free to use and "
    "edit on any project."
)

# Shared Part 1 and Part 3 clauses ------------------------------------------------------------

def submittals(extra_action, extra_info):
    return [
        ("h2", "1.4 SUBMITTALS"),
        ("p", "A. Action submittals:"),
        ("s", "1. Product data for each item, showing ratings, dimensions, weights and options."),
        *[("s", f"{i + 2}. {t}") for i, t in enumerate(extra_action)],
        ("p", "B. Informational submittals:"),
        ("s", "1. Certification: CSA or cUL listing, or a field evaluation arranged before shipment."),
        *[("s", f"{i + 2}. {t}") for i, t in enumerate(extra_info)],
        ("p", "C. Closeout submittals: operation and maintenance manuals, as-built drawings, certified factory test reports and recommended spare parts list, in [PDF] [and] [printed] form."),
    ]


def delivery(item):
    return [
        ("h2", "1.6 DELIVERY, STORAGE AND HANDLING"),
        ("p", f"A. Deliver {item} with protective packaging, desiccant and impact recorders where specified. Record recorder readings on arrival."),
        ("p", "B. Store indoors in a clean, dry, heated space, or outdoors under cover, off the ground, with space heaters energized from a temporary supply."),
        ("p", "C. Lift and move only at the marked lifting points, following the manufacturer's instructions."),
    ]


def warranty(months="[24]"):
    return [
        ("h2", "1.7 WARRANTY"),
        ("p", f"A. Manufacturer's warranty against defects in materials and workmanship for {months} months from substantial performance, or [30] months from shipment, whichever comes first."),
    ]


def qa(extra):
    return [
        ("h2", "1.5 QUALITY ASSURANCE"),
        ("p", "A. Manufacturer: regularly produces equipment of the type and ratings specified, with a documented quality management system."),
        ("p", "B. Equipment shall be certified to the applicable CSA standards and bear the certification mark, or be field-evaluated and labelled before energization to the satisfaction of the authority having jurisdiction."),
        ("note", "Optional domestic-content clause. Keep it only if the owner requires Canadian-made equipment."),
        ("p", "C. [Equipment shall be manufactured and factory-tested in Canada.]"),
        *extra,
    ]


def manufacturers():
    return [
        ("h2", "2.1 MANUFACTURERS"),
        ("p", "A. Acceptable manufacturers or suppliers: [list], or approved equal submitted [10] working days before tender close."),
        ("note", "List at least two acceptable sources where the owner's procurement rules require competition."),
    ]


def execution(item, neta_tests, training_hours="[4]"):
    return [
        ("h1", "PART 3 EXECUTION"),
        ("h2", "3.1 EXAMINATION"),
        ("p", f"A. Verify pads, anchors, clearances, cable entries and grounding are ready before setting the {item}. Report defects in writing before proceeding."),
        ("h2", "3.2 INSTALLATION"),
        ("p", f"A. Install the {item} to the manufacturer's instructions, CSA C22.1 (Canadian Electrical Code, Part I) and the requirements of the authority having jurisdiction."),
        ("p", "B. Anchor to the foundation; level and align. Torque bolted connections to the manufacturer's values and record them."),
        ("p", "C. Ground to the station or building grounding system as shown."),
        ("h2", "3.3 FIELD QUALITY CONTROL"),
        ("p", "A. Perform acceptance tests to ANSI/NETA ATS with a qualified testing agency, including:"),
        *[("s", f"{i + 1}. {t}") for i, t in enumerate(neta_tests)],
        ("p", "B. Submit test reports before energization. Correct deficiencies and retest."),
        ("h2", "3.4 COMMISSIONING"),
        ("p", "A. Provide the manufacturer's or supplier's field representative for [start-up] [and] [energization] [and] [protection setting verification]."),
        ("h2", "3.5 TRAINING"),
        ("p", f"A. Provide {training_hours} hours of training for the owner's operating staff on operation, safety and routine maintenance."),
        ("p", "END OF SECTION", "end"),
    ]


# Sections --------------------------------------------------------------------------------------

SECTIONS = [
    {
        "file": "26-12-19-pad-mounted-liquid-filled-transformers.docx",
        "number": "26 12 19",
        "title": "PAD-MOUNTED, LIQUID-FILLED, MEDIUM-VOLTAGE TRANSFORMERS",
        "body": [
            ("h1", "PART 1 GENERAL"),
            ("h2", "1.1 SUMMARY"),
            ("p", "A. This section covers three-phase, pad-mounted, liquid-filled distribution transformers with compartmented, tamper-resistant enclosures, [loop] [radial] feed, rated [75] to [5,000] kVA."),
            ("h2", "1.2 RELATED SECTIONS"),
            ("p", "A. Section 26 05 26 Grounding and Bonding. Section 26 13 xx Medium-Voltage Switchgear. Section 03 30 00 Cast-in-Place Concrete (pads)."),
            ("h2", "1.3 REFERENCES"),
            ("p", "A. CSA C2.1, Single-phase and three-phase liquid-filled distribution transformers."),
            ("p", "B. CSA C227.4, Three-phase, pad-mounted distribution transformers with separable insulated high-voltage connectors."),
            ("p", "C. CSA C802.1, Minimum efficiency values for liquid-filled distribution transformers."),
            ("p", "D. IEEE C57.12.00, C57.12.34, C57.12.90 and C57.12.28 (enclosure integrity) [and C57.12.29 (coastal environments)]."),
            ("p", "E. ASTM D3487 (mineral insulating oil) or ASTM D6871 (natural ester fluid)."),
            ("p", "F. Requirements of the local distribution company where the transformer is utility-owned or utility-interconnected."),
            *submittals(["Outline drawing with pad layout, cable entry window and bushing positions.", "Nameplate drawing and connection diagram.", "Calculated no-load and load losses, and efficiency at the CSA C802.1 reference load."], ["Certified test reports.", "Fluid data sheet."]),
            *qa([]),
            *delivery("transformers"),
            *warranty(),
            ("h1", "PART 2 PRODUCTS"),
            *manufacturers(),
            ("h2", "2.2 RATINGS"),
            ("p", "A. Rating: [____] kVA, [ONAN] [ONAN/ONAF], 65 °C average winding rise."),
            ("p", "B. Primary: [27.6] [13.8] [4.16] kV [delta] [grounded wye], BIL [125] [95] [60] kV. Taps: four 2.5% full-capacity taps, two above and two below rated voltage, de-energized tap changer."),
            ("p", "C. Secondary: [600Y/347] [480Y/277] [208Y/120] V, BIL [30] kV."),
            ("p", "D. Impedance: [5.75]% ± 7.5% tolerance, or as required by the fault study."),
            ("p", "E. Efficiency: not less than the CSA C802.1 minimum for the rating."),
            ("p", "F. Sound level: not above the NEMA TR 1 value for the rating."),
            ("h2", "2.3 CONSTRUCTION"),
            ("p", "A. Core: grain-oriented electrical steel, step-lap mitred joints."),
            ("p", "B. Windings: [copper] [aluminum], with thermally upgraded paper insulation."),
            ("p", "C. Fluid: [mineral oil to ASTM D3487] [natural ester fluid to ASTM D6871, less-flammable listed]."),
            ("note", "Natural ester fluid improves fire safety and paper life but has a higher pour point. Check the minimum site temperature."),
            ("p", "D. Tank: sealed, welded steel, designed for 7 psig without permanent deformation. Pressure relief device."),
            ("p", "E. High-voltage compartment: [dead-front with 200 A load-break bushing wells and inserts] [live-front], [loop feed] [radial feed], with [two-position] [four-position] load-break oil switch."),
            ("p", "F. Overcurrent protection: Bay-O-Net expulsion fuses in series with partial-range current-limiting fuses, or as required by the utility."),
            ("p", "G. Surge arresters: [elbow] [distribution-class] arresters on each primary phase."),
            ("p", "H. Low-voltage compartment: spade terminals sized for [____] conductors per phase, with neutral bushing and removable ground strap."),
            ("h2", "2.4 ACCESSORIES"),
            ("p", "A. Liquid-level gauge, dial-type thermometer with maximum pointer, pressure-vacuum gauge, drain valve with sampling device, and fill plug."),
            ("p", "B. [Alarm contacts on gauges wired to a terminal block.] [Provision for remote monitoring.]"),
            ("h2", "2.5 ENCLOSURE AND FINISH"),
            ("p", "A. Tamper-resistant, padlockable compartments to IEEE C57.12.28. Finish: [Munsell 7GY3.29/1.5 green] [owner's colour]."),
            ("p", "B. [Stainless steel cabinet and base for coastal or corrosive sites.]"),
            ("h2", "2.6 FACTORY TESTS"),
            ("p", "A. Routine tests to IEEE C57.12.90 and CSA C2.1, including resistance, ratio, polarity and phase relation, no-load and load losses, excitation current, impedance, applied and induced potential, and leak test."),
            ("p", "B. [Witnessed factory test: give [10] working days' notice.]"),
            *execution("transformer", ["Visual and mechanical inspection.", "Insulation resistance of windings.", "Turns ratio on all tap positions.", "Winding resistance.", "Fluid sample for dielectric breakdown, moisture and dissolved gas analysis (baseline)."]),
        ],
    },
    {
        "file": "26-13-13-medium-voltage-metal-clad-switchgear.docx",
        "number": "26 13 13",
        "title": "MEDIUM-VOLTAGE CIRCUIT BREAKER SWITCHGEAR (METAL-CLAD)",
        "body": [
            ("h1", "PART 1 GENERAL"),
            ("h2", "1.1 SUMMARY"),
            ("p", "A. This section covers indoor [and outdoor walk-in] [and outdoor non-walk-in] metal-clad switchgear with draw-out vacuum circuit breakers, rated [5] [15] [27] [38] kV."),
            ("h2", "1.2 RELATED SECTIONS"),
            ("p", "A. Section 26 05 26 Grounding and Bonding. Section 26 05 73 Power System Studies. Section 26 12 xx Medium-Voltage Transformers."),
            ("h2", "1.3 REFERENCES"),
            ("p", "A. CSA C22.2 No. 31, Switchgear assemblies."),
            ("p", "B. IEEE C37.20.2, Metal-clad switchgear. [IEEE C37.20.7, Testing of switchgear for internal arcing faults.]"),
            ("p", "C. IEEE C37.04, C37.06 and C37.09, AC high-voltage circuit breakers."),
            ("p", "D. IEEE C37.90, Relays and relay systems. CSA C22.1, Canadian Electrical Code, Part I."),
            *submittals(["Shop drawings: front view, plan, sections, single-line, three-line and schematic diagrams.", "Bill of materials with device ratings.", "Relay and meter list with communication protocol."], ["Certified design (type) test reports for the breaker and the assembly.", "Arc-resistance test report, if arc-resistant gear is specified."]),
            *qa([("p", "D. The switchgear assembly, breakers and protective relays shall be from one manufacturer's coordinated design, with one source of responsibility.")]),
            *delivery("switchgear"),
            *warranty(),
            ("h1", "PART 2 PRODUCTS"),
            *manufacturers(),
            ("h2", "2.2 RATINGS"),
            ("p", "A. System: [13.8] kV, three-phase, [three] [four]-wire, 60 Hz, [solidly] [resistance] grounded."),
            ("p", "B. Rated maximum voltage [15] kV; BIL [95] kV; main bus [1,200] [2,000] [3,000] A continuous."),
            ("p", "C. Short-circuit rating: [25] [40] [50] kA symmetrical, with momentary and short-time ratings to IEEE C37.06."),
            ("p", "D. [Arc-resistant: Type [2B], [25] kA for [0.5] s, with plenum venting to [outdoors] [a safe location shown].]"),
            ("note", "Arc-resistant construction reduces risk to operators; confirm room height and plenum route early, as they drive the layout."),
            ("h2", "2.3 CONSTRUCTION"),
            ("p", "A. Grounded steel compartments separating each breaker, main bus, cable connections and low-voltage instruments. Automatic shutters over primary stabs."),
            ("p", "B. Bus: [silver-plated copper], fully insulated, with insulated joint boots. Ground bus: copper, full length."),
            ("p", "C. Breakers: draw-out vacuum, electrically operated, stored-energy mechanism, with connected, test and disconnected positions. [Remote racking device.]"),
            ("p", "D. Interlocks preventing racking a closed breaker and preventing insertion of a breaker of the wrong rating."),
            ("p", "E. Instrument transformers: CTs of accuracy class to suit the relays and meters; VTs with draw-out primary fuses."),
            ("p", "F. Protection: multifunction microprocessor relays providing [50/51, 50N/51N, 27/59, 81, 86] as shown, with [DNP3] [Modbus TCP] communication."),
            ("note", "Choose the protocol that matches the owner's SCADA system."),
            ("p", "G. Control power: [125] V DC from a [battery and charger specified elsewhere] [battery and charger included]."),
            ("p", "H. Space heaters with thermostat in every outdoor compartment and in indoor cable compartments."),
            ("h2", "2.4 FACTORY TESTS"),
            ("p", "A. Production tests to IEEE C37.20.2 and CSA C22.2 No. 31: power-frequency withstand, mechanical operation, wiring and control checks, and relay injection with the approved settings."),
            ("p", "B. [Witnessed factory acceptance test: give [10] working days' notice.]"),
            *execution("switchgear", ["Visual and mechanical inspection, including interlocks and racking.", "Insulation resistance of bus and breakers.", "Contact resistance of bus joints and breaker poles.", "Vacuum bottle integrity (overpotential) test.", "Breaker timing and trip tests.", "Secondary injection of protective relays with final settings.", "CT ratio and polarity, VT ratio."], "[8]"),
        ],
    },
    {
        "file": "26-22-13-low-voltage-dry-type-transformers.docx",
        "number": "26 22 13",
        "title": "LOW-VOLTAGE DISTRIBUTION TRANSFORMERS (DRY-TYPE)",
        "body": [
            ("h1", "PART 1 GENERAL"),
            ("h2", "1.1 SUMMARY"),
            ("p", "A. This section covers ventilated, dry-type, general-purpose, [K-factor] [harmonic-mitigating] transformers rated 600 V class and below, single- and three-phase, [15] to [1,000] kVA."),
            ("h2", "1.2 RELATED SECTIONS"),
            ("p", "A. Section 26 05 26 Grounding and Bonding. Section 26 24 13 Switchboards. Section 26 24 16 Panelboards."),
            ("h2", "1.3 REFERENCES"),
            ("p", "A. CSA C9, Dry-type transformers."),
            ("p", "B. CSA C22.2 No. 47, Air-cooled transformers (dry type)."),
            ("p", "C. CSA C802.2, Minimum efficiency values for dry-type transformers, and the federal Energy Efficiency Regulations."),
            ("p", "D. IEEE C57.12.01, C57.12.91 and C57.110 (non-sinusoidal load currents). NEMA ST 20."),
            *submittals(["Outline drawing with mounting, ventilation clearances and weight.", "Wiring diagram and nameplate data.", "No-load and load losses, and efficiency at the CSA C802.2 reference load."], ["Sound level data."]),
            *qa([]),
            *delivery("transformers"),
            *warranty(),
            ("h1", "PART 2 PRODUCTS"),
            *manufacturers(),
            ("h2", "2.2 RATINGS"),
            ("p", "A. Primary [600] [480] V delta; secondary [208Y/120] [600Y/347] [240/120] V; 60 Hz."),
            ("p", "B. Taps: six 2.5% full-capacity taps, two above and four below rated voltage, for units 15 kVA and larger."),
            ("p", "C. Insulation system: 220 °C class, with [150] [115] [80] °C average winding rise."),
            ("note", "A lower temperature rise lowers losses and extends life, at a higher first cost and weight."),
            ("p", "D. [K-factor rating: K-[4] [9] [13] [20] to IEEE C57.110, with a neutral rated 200% of phase current.]"),
            ("p", "E. Efficiency: not less than the CSA C802.2 minimum for the rating."),
            ("p", "F. Sound level: not above NEMA ST 20 values [less 3 dB]."),
            ("h2", "2.3 CONSTRUCTION"),
            ("p", "A. Core: grain-oriented electrical steel, mitred joints, clamped to limit noise."),
            ("p", "B. Windings: [copper] [aluminum], vacuum-pressure impregnated with varnish."),
            ("p", "C. [Electrostatic shield between primary and secondary, grounded to the enclosure.]"),
            ("p", "D. Vibration isolation pads between the core and coil and the enclosure."),
            ("p", "E. Enclosure: ventilated [CSA Type 2 (indoor drip-proof)] [CSA Type 3R (outdoor)], with removable front and rear panels, finished in [ASA 61 grey]. Maximum enclosure temperature rise as required by CSA C22.2 No. 47."),
            ("p", "F. [Wall-mounting brackets for units [75] kVA and below.]"),
            ("h2", "2.4 FACTORY TESTS"),
            ("p", "A. Routine tests to IEEE C57.12.91 and CSA C9: resistance, ratio, polarity and phase relation, no-load loss, excitation current, impedance and load loss, applied and induced potential."),
            *execution("transformer", ["Visual and mechanical inspection, including removal of shipping braces.", "Insulation resistance of windings.", "Turns ratio on the as-left tap.", "Secondary voltage under load; adjust taps as needed."], "[2]"),
        ],
    },
    {
        "file": "26-24-13-switchboards.docx",
        "number": "26 24 13",
        "title": "SWITCHBOARDS",
        "body": [
            ("h1", "PART 1 GENERAL"),
            ("h2", "1.1 SUMMARY"),
            ("p", "A. This section covers [service-entrance] [distribution] switchboards rated 600 V and below, [400] to [4,000] A, with [moulded-case] [insulated-case] [power] circuit breakers [and utility metering]."),
            ("h2", "1.2 RELATED SECTIONS"),
            ("p", "A. Section 26 05 26 Grounding and Bonding. Section 26 05 73 Power System Studies. Section 26 22 13 Low-Voltage Distribution Transformers. Section 26 36 00 Transfer Switches."),
            ("h2", "1.3 REFERENCES"),
            ("p", "A. CSA C22.2 No. 244, Switchboards (binational with UL 891)."),
            ("p", "B. CSA C22.2 No. 5, Moulded-case circuit breakers."),
            ("p", "C. NEMA PB 2. CSA C22.1, Canadian Electrical Code, Part I."),
            ("p", "D. Measurement Canada requirements and the local utility's metering standard, where utility metering is included."),
            *submittals(["Shop drawings: front elevation, plan, single-line, bus ratings and conduit areas.", "Breaker list with frame, sensor and trip unit settings range.", "Utility metering compartment drawing approved by the utility, where applicable."], ["Short-circuit and coordination study results applied to trip settings (by others or included)."]),
            *qa([]),
            *delivery("switchboards"),
            *warranty(),
            ("h1", "PART 2 PRODUCTS"),
            *manufacturers(),
            ("h2", "2.2 RATINGS"),
            ("p", "A. System: [600Y/347] [208Y/120] V, three-phase, four-wire, 60 Hz."),
            ("p", "B. Main bus: [____] A continuous, [copper, tin-plated] [copper, silver-plated]. Neutral: 100% rated. Ground bus: full length."),
            ("p", "C. Short-circuit rating: [65] [100] kA symmetrical, fully rated. Series ratings are not acceptable [unless approved]."),
            ("p", "D. Service-entrance labelled where used as service equipment, with ground-fault protection on mains rated 1,000 A or more where required by the Canadian Electrical Code."),
            ("h2", "2.3 CONSTRUCTION"),
            ("p", "A. Free-standing, front-accessible [front and rear accessible] sections, [CSA Type 1] [CSA Type 3R outdoor]."),
            ("p", "B. Main device: [fixed] [draw-out] [insulated-case] [power] circuit breaker with electronic trip unit, LSI[G] functions [and arc-energy reduction maintenance switch]."),
            ("note", "An arc-energy reduction switch lowers incident energy during maintenance; most owners now ask for it on mains 1,200 A and above."),
            ("p", "C. Feeders: [group-mounted moulded-case] [individually mounted] breakers, as shown on the single-line."),
            ("p", "D. Metering: multifunction power meter on each main with [Modbus TCP] [BACnet] communication. [Utility metering compartment to the utility's standard, with sealable doors.]"),
            ("p", "E. Surge protective device on each main bus: [Type 1] [Type 2], [____] kA per mode."),
            ("p", "F. [Integrated transformers and panelboards in ventilated sections, where shown.]"),
            ("p", "G. Laminated nameplates, mechanically fastened, for each section and device."),
            ("h2", "2.4 FACTORY TESTS"),
            ("p", "A. Production tests to CSA C22.2 No. 244: dielectric, continuity, polarity and mechanical operation of devices; functional checks of controls and meters."),
            ("p", "B. [Witnessed factory acceptance test: give [10] working days' notice.]"),
            *execution("switchboard", ["Visual and mechanical inspection, including bolted bus joints.", "Insulation resistance of bus.", "Contact resistance of main bus joints.", "Primary or secondary injection of main breaker trip units with final settings.", "Ground-fault protection performance test, where provided."], "[2]"),
        ],
    },
]


# Rendering -------------------------------------------------------------------------------------

def shade(paragraph, hex_fill):
    ppr = paragraph._p.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_fill)
    ppr.append(shd)


def page_field(run):
    for kind, text in (("begin", None), (None, "PAGE"), ("end", None)):
        if kind:
            el = OxmlElement("w:fldChar")
            el.set(qn("w:fldCharType"), kind)
        else:
            el = OxmlElement("w:instrText")
            el.set(qn("xml:space"), "preserve")
            el.text = text
        run._r.append(el)


def build(sec):
    doc = Document()
    section = doc.sections[0]
    section.page_width, section.page_height = Cm(21.59), Cm(27.94)
    for side in ("left_margin", "right_margin"):
        setattr(section, side, Cm(2.2))
    section.top_margin, section.bottom_margin = Cm(2.0), Cm(2.0)

    normal = doc.styles["Normal"]
    normal.font.name = "Arial"
    normal.font.size = Pt(10)
    normal.paragraph_format.space_after = Pt(4)

    head = section.header.paragraphs[0]
    head.text = f"Section {sec['number']} · {sec['title'].title()}"
    head.runs[0].font.size = Pt(8)
    head.runs[0].font.color.rgb = GREY
    foot = section.footer.paragraphs[0]
    foot.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = foot.add_run(f"{sec['number']} – ")
    r.font.size = Pt(8)
    pr = foot.add_run()
    pr.font.size = Pt(8)
    page_field(pr)
    r2 = foot.add_run("   ·   Guide specification from CANDRON Energy Inc., providers of Canadian-made electrical equipment · candron.com")
    r2.font.size = Pt(8)
    r2.font.color.rgb = GREY

    t = doc.add_paragraph()
    run = t.add_run(f"SECTION {sec['number']}")
    run.bold = True
    run.font.size = Pt(12)
    run.font.color.rgb = BLUE
    t2 = doc.add_paragraph()
    run = t2.add_run(sec["title"])
    run.bold = True
    run.font.size = Pt(14)
    run.font.color.rgb = NAVY
    intro = doc.add_paragraph()
    ir = intro.add_run(COMMON_INTRO)
    ir.italic = True
    ir.font.size = Pt(9)
    ir.font.color.rgb = GREY
    shade(intro, "EEF2FA")

    for item in sec["body"]:
        kind, text = item[0], item[1]
        if kind == "h1":
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(14)
            rr = p.add_run(text)
            rr.bold = True
            rr.font.size = Pt(11)
            rr.font.color.rgb = NAVY
        elif kind == "h2":
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(8)
            p.paragraph_format.keep_with_next = True
            rr = p.add_run(text)
            rr.bold = True
        elif kind == "p":
            p = doc.add_paragraph(text)
            p.paragraph_format.left_indent = Cm(0.6)
            if len(item) > 2 and item[2] == "end":
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                p.paragraph_format.left_indent = Cm(0)
                p.paragraph_format.space_before = Pt(18)
                p.runs[0].bold = True
        elif kind == "s":
            p = doc.add_paragraph(text)
            p.paragraph_format.left_indent = Cm(1.4)
        elif kind == "note":
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Cm(0.6)
            rr = p.add_run(f"Specifier note: {text}")
            rr.italic = True
            rr.font.size = Pt(9)
            rr.font.color.rgb = BLUE
            shade(p, "EEF2FA")

    doc.core_properties.title = f"Section {sec['number']} {sec['title'].title()}"
    doc.core_properties.author = "CANDRON Energy Inc."
    OUT.mkdir(parents=True, exist_ok=True)
    doc.save(OUT / sec["file"])
    print("wrote", OUT / sec["file"])


if __name__ == "__main__":
    for s in SECTIONS:
        build(s)
